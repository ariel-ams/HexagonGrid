from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs" / "HONEYCOMB_CORE_EXTRACTION_BRIEF.md"
DEFAULT_OUTPUT = ROOT / "docs" / "Honeycomb-Wayfinder-Core-Extraction-Brief.docx"
PAGE_BREAK_HEADINGS = {
    "Generalization methodology",
    "Creating a similar game",
}


def set_font(run, name: str, size: float, color: str = "000000", bold: bool | None = None, italic: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_style(style, size: float, color: str, before: float, after: float, line_spacing: float, bold: bool = False):
    style.font.name = "Arial"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_style(doc.styles["Normal"], 11, "000000", 0, 8, 1.15)
    configure_style(doc.styles["Heading 1"], 20, "000000", 20, 6, 1.0)
    configure_style(doc.styles["Heading 2"], 16, "000000", 18, 6, 1.0)
    configure_style(doc.styles["Heading 3"], 14, "434343", 16, 4, 1.0)

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        configure_style(style, 11, "000000", 0, 4, 1.15)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def add_inline(paragraph, text: str, default_size: float = 11):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[1:-1] if is_code else part[2:-2] if is_bold else part
        run = paragraph.add_run(value)
        set_font(
            run,
            "Courier New" if is_code else "Arial",
            9.5 if is_code else default_size,
            bold=True if is_bold else None,
        )


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def set_picture_alt(inline_shape, alt: str):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", alt)
    doc_pr.set("descr", alt)


def create_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "decimal")
    level.append(num_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    paragraph_properties.append(indent)
    level.append(paragraph_properties)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num_properties.append(level)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_properties.append(num)
    paragraph_properties.append(num_properties)


def add_code_block(doc: Document, lines: list[str]):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.16)
    paragraph.paragraph_format.right_indent = Inches(0.16)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    shade_paragraph(paragraph, "F5F5F5")
    run = paragraph.add_run("\n".join(lines))
    set_font(run, "Courier New", 8.3, "202124")


def add_image(doc: Document, source_dir: Path, alt: str, relative_path: str):
    path = (source_dir / relative_path).resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(6.3))
    set_picture_alt(shape, alt)


def build(source: Path, output: Path):
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Honeycomb Wayfinder: Core Extraction and Adaptation Methodology"
    doc.core_properties.subject = "Reusable tactical hex-engine extraction, theme packaging, and content authoring methodology"
    doc.core_properties.keywords = "Honeycomb Wayfinder, hex grid, tactical engine, extraction, methodology, theme packs, content tools"

    lines = source.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    title_written = False
    numbering_id: int | None = None

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line:
            numbering_id = None
            continue

        image_match = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line)
        if image_match:
            add_image(doc, source.parent, image_match.group(1), image_match.group(2))
            continue

        if line.startswith("# "):
            numbering_id = None
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(line[2:])
            set_font(run, "Arial", 26, "000000", bold=False)
            title_written = True
            continue

        if line.startswith("## "):
            numbering_id = None
            heading = line[3:]
            if title_written and heading in PAGE_BREAK_HEADINGS:
                doc.add_page_break()
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, heading, 20)
            continue

        if line.startswith("### "):
            numbering_id = None
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, line[4:], 16)
            continue

        if line.startswith("- "):
            numbering_id = None
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            continue

        number_match = re.match(r"^\d+\.\s+(.*)$", line)
        if number_match:
            if numbering_id is None:
                numbering_id = create_numbering_id(doc)
            paragraph = doc.add_paragraph(style="List Number")
            apply_numbering(paragraph, numbering_id)
            add_inline(paragraph, number_match.group(1))
            continue

        if line.startswith("*") and line.endswith("*") and len(line) > 2:
            numbering_id = None
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run(line[1:-1])
            set_font(run, "Arial", 9, "555555", italic=True)
            continue

        numbering_id = None
        paragraph = doc.add_paragraph()
        add_inline(paragraph, line)

    if in_code:
        add_code_block(doc, code_lines)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
